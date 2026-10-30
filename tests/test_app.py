"""The real page in Edge: from folders or tables to the plots and files. All data here is synthetic."""
import csv
import datetime
import http.server
import io
import math
import os
import re
import shutil
import struct
import threading
import zipfile
from functools import partial
from pathlib import Path

import numpy as np
import openpyxl
import pytest
import statsmodels.api as sm
import xlwt

APP = Path(__file__).resolve().parents[1] / "app"


def make_clips(root, n=120, b0=-1.0, b1=1.6, seed=1, site="SITE01", year=2024, seconds=3.0):
    """Empty files named the way BirdNET names its clips: only the names matter. Returns the scores and verdicts."""
    rng = np.random.default_rng(seed)
    (root / "Positive").mkdir(parents=True, exist_ok=True)
    (root / "Negative").mkdir(parents=True, exist_ok=True)
    scores, verdicts = [], []
    for i in range(n):
        s = round(float(rng.uniform(0.1, 1.0)), 3)
        correct = rng.random() < 1 / (1 + math.exp(-(b0 + b1 * math.log(max(s, 1e-6) / max(1 - s, 1e-6)))))
        start = float(rng.integers(0, 3000))
        name = f"{s:.3f}_{i + 1}_{site}_{year}05{1 + i % 28:02d}_2030{i % 60:02d}_{start:.1f}s_{start + seconds:.1f}s.wav"
        (root / ("Positive" if correct else "Negative") / name).touch()
        scores.append(s)
        verdicts.append(int(correct))
    return np.array(scores), np.array(verdicts)


def browse(page, row, kind, folder):
    with page.expect_file_chooser() as chooser:
        page.locator(".species-row").nth(row).locator(f'.slot[data-part="{kind}"] .browse').click()
    chooser.value.set_files(str(folder))


def add_species(page, row, name, root):
    if row > 0:
        page.click("#add-species")
    page.locator(".species-name").nth(row).fill(name)
    browse(page, row, "pos", root / "Positive")
    browse(page, row, "neg", root / "Negative")


def download(page, selector):
    with page.expect_download() as d:
        page.click(selector)
    return Path(d.value.path()).read_bytes(), d.value.suggested_filename


def read_csv(raw, delimiter=","):
    return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig")), delimiter=delimiter))


def svg_of(page, i=0):
    return page.locator(".plot-canvas svg").nth(i).evaluate("(s) => s.outerHTML")


def threshold(scores, verdicts, p, logit=True):
    x = np.log(scores / (1 - np.clip(scores, 0, 1 - 1e-6))) if logit else scores
    b0, b1 = sm.GLM(verdicts, sm.add_constant(x), family=sm.families.Binomial()).fit(tol=1e-12).params
    q = (math.log(p / (1 - p)) - b0) / b1
    return 1 / (1 + math.exp(-q)) if logit else q


@pytest.fixture
def one_species(page, tmp_path):
    scores, verdicts = make_clips(tmp_path / "a")
    add_species(page, 0, "Turdus_merula", tmp_path / "a")
    return scores, verdicts


# ── Folders to plot and files ────────────────────────────────────────────────────────────────────────────
def test_folders_give_the_csv_and_the_right_threshold(page, one_species):
    scores, verdicts = one_species
    state = page.locator('.slot[data-part="pos"] .slot-count').inner_text()
    assert state.startswith(f"{verdicts.sum()} clips")
    page.click("#generate")
    page.wait_for_selector(".plot-canvas svg")
    raw, name = download(page, "#download-data")
    rows = read_csv(raw)
    assert name == "birdscore_detections.csv" and len(rows) == len(scores)
    assert sorted(float(r["confidence_score"]) for r in rows) == sorted(scores.tolist())
    assert sum(int(r["correct"]) for r in rows) == verdicts.sum()
    r = rows[0]
    assert r["species"] == "Turdus_merula" and r["result"] in ("positive", "negative")
    assert r["site_or_recorder"] == "SITE01" and re.fullmatch(r"2024-05-\d\d", r["date"]) and re.fullmatch(r"20:30:\d\d", r["time"])
    assert float(r["clip_end_s"]) - float(r["clip_start_s"]) == pytest.approx(3.0)
    assert r["recording"].startswith("SITE01_2024") and r["file_name"].endswith(".wav") and r["folder"] == "Positive" or r["folder"] == "Negative"

    # The default mark is p = 0.90, and the table shows the fit that statsmodels finds
    cell = page.locator("#results-table tbody td").nth(4).inner_text()
    assert cell.startswith(f"{threshold(scores, verdicts, 0.9):.2f} (")
    results = read_csv(download(page, "#download-results")[0])
    assert float(results[0]["score_at_p_0.9"]) == pytest.approx(threshold(scores, verdicts, 0.9), abs=1e-3)
    assert float(results[0]["score_at_p_0.9_ci_low"]) < float(results[0]["score_at_p_0.9"]) < float(results[0]["score_at_p_0.9_ci_high"])
    assert results[0]["score_scale"] == "logit"


def test_the_plot_has_the_defaults_asked_for(page, one_species):
    page.click("#generate")
    svg = svg_of(page)
    assert "Turdus merula" in svg and "font-style=\"italic\"" in svg       # the title is the species name
    assert ">Confidence score<" in svg and ">True positive rate<" in svg
    assert re.search(r">\d+ positive · \d+ negative<", svg)


def test_settings_change_the_plot_without_generating_again(page, one_species):
    page.click("#generate")
    page.locator("#title-list input").nth(0).fill("Common blackbird")
    page.locator("#title-list input").nth(1).fill("Validated in spring")
    page.fill("#x-label", "BirdNET score")
    page.fill("#y-label", "Precision")
    page.uncheck("#italic")
    page.wait_for_function("!document.querySelector('.plot-canvas svg').outerHTML.includes('font-style')")     # the last change
    svg = svg_of(page)
    assert all(text in svg for text in ["Common blackbird", "Validated in spring", ">BirdNET score<", ">Precision<"])


def test_thresholds_by_probability_and_by_score(page, one_species):
    scores, verdicts = one_species
    page.click("#generate")
    page.click('#target-chips [data-value="0.95"]')
    page.fill("#target-new", "0,97")
    page.click("#target-add")
    page.fill("#score-new", "0.5")
    page.click("#score-add")
    page.wait_for_function("document.querySelectorAll('#results-table th').length === 8")
    heads = page.locator("#results-table th").all_inner_texts()
    assert heads[4:] == ["Score for p = 0.90", "Score for p = 0.95", "Score for p = 0.97", "p at score 0.50"]
    assert page.locator("#results-table tbody td").nth(6).inner_text().startswith(f"{threshold(scores, verdicts, 0.97):.2f}")
    dashed = page.locator('.plot-canvas svg path[stroke-dasharray]')
    assert dashed.evaluate_all("(p) => p.filter((x) => x.getAttribute('d').includes('V')).length") == 4     # the fifth is the legend's key
    page.click('#target-chips [data-value="0.9"]')
    page.click('#target-chips [data-value="0.95"]')
    page.click('#target-chips [data-value="0.97"]')
    page.click('#score-chips [data-value="0.5"]')
    page.wait_for_function("document.querySelectorAll('.plot-canvas svg path[stroke-dasharray]').length === 0")
    assert page.locator('#target-chips [data-value="0.97"]').count() == 0       # a custom value goes when turned off
    assert page.locator('#target-chips [data-value="0.95"]').count() == 1       # a preset stays


def test_bad_threshold_values_say_how_to_fix_them(page):
    page.fill("#target-new", "1.5")
    page.click("#target-add")
    assert page.is_visible("#target-error") and "above 0 and below 1" in page.inner_text("#target-error")
    page.fill("#target-new", "0.8")
    page.click("#target-add")
    assert page.is_hidden("#target-error")
    assert page.locator('#target-chips [data-value="0.8"]').get_attribute("aria-pressed") == "true"


def test_jitter(page, one_species):
    cys = lambda: page.locator(".plot-canvas svg g[fill='#0d9488'] circle").evaluate_all("(c) => c.map((x) => x.getAttribute('cy'))")
    cxs = lambda: page.locator(".plot-canvas svg g[fill='#0d9488'] circle").evaluate_all("(c) => c.map((x) => x.getAttribute('cx'))")
    page.click("#generate")
    spread, xs = cys(), cxs()
    assert len(set(spread)) > 5
    page.fill("#x-label", "again")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('again')")
    assert cys() == spread                                    # the same points on every drawing
    page.uncheck("#jitter")
    page.wait_for_function("new Set([...document.querySelectorAll(\".plot-canvas svg g[fill='#0d9488'] circle\")].map((c) => c.getAttribute('cy'))).size === 1")
    assert cxs() == xs                                        # jitter never moves a score


def narrow_clips(root):
    """Scores only from 0.30 to 0.80, with a slow rise, so that high targets lie beyond the scores checked."""
    (root / "Positive").mkdir(parents=True)
    (root / "Negative").mkdir(parents=True)
    for i in range(50):
        s = 0.3 + i / 100
        (root / ("Positive" if i % 5 else "Negative") / f"{s:.3f}_{i}_r_{i}.0s_{i + 3}.0s.wav").touch()
    (root / "Positive" / "0.800_99_r_0.0s_3.0s.wav").touch()


def test_the_curve_stops_at_the_scores_checked(page, tmp_path):
    narrow_clips(tmp_path / "a")
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.click('#target-chips [data-value="0.9"]')
    page.click('#target-chips [data-value="0.99"]')
    page.click("#generate")
    svg = svg_of(page)
    assert "Extrapolated" not in svg
    curve = page.locator('.plot-canvas svg path[stroke-width="2"][fill="none"]').first.get_attribute("d")
    xs = [float(x) for x in re.findall(r"[ML]([\d.]+) ", curve)]
    left, width = 64, 660
    assert min(xs) == pytest.approx(left + 0.30 * width, abs=0.2) and max(xs) == pytest.approx(left + 0.80 * width, abs=0.2)
    # The threshold beyond them is not drawn, but it stays in the table with a note
    assert page.locator(".plot-canvas svg path[stroke-dasharray]").count() == 0 and ">Threshold<" not in svg
    assert "is not drawn" in page.inner_text("#notes")
    assert re.match(r"\d\.\d\d", page.locator("#results-table tbody td").nth(4).inner_text())


def test_hovering_reads_the_curve_only_where_it_was_checked(page, tmp_path):
    narrow_clips(tmp_path / "a")
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.click("#generate")
    svg = page.locator(".plot-canvas svg")
    svg.scroll_into_view_if_needed()
    to_screen = "(s, [x, y]) => { const p = new DOMPoint(x, y).matrixTransform(s.getScreenCTM()); return [p.x, p.y]; }"
    inside = svg.evaluate(to_screen, [64 + 0.55 * 660, 300])
    outside = svg.evaluate(to_screen, [64 + 0.10 * 660, 300])
    page.mouse.move(*inside)
    page.wait_for_selector("#tip:not([hidden])")
    assert page.inner_text("#tip").startswith("p = ")
    page.mouse.move(*outside)
    page.wait_for_selector("#tip", state="hidden")


def test_a_threshold_is_only_a_vertical_line_with_its_score(page, one_species):
    scores, verdicts = one_species
    page.click("#generate")
    lines = page.locator(".plot-canvas svg path[stroke-dasharray]").evaluate_all("(p) => p.map((x) => x.getAttribute('d'))")
    drawn = [d for d in lines if "V" in d]
    assert len(drawn) == 1 and "H" not in drawn[0]                             # the other dashed path is the legend's key
    svg = svg_of(page)
    assert page.locator('.plot-canvas svg circle[r="4"][stroke-width="2"]').count() == 0     # no dot on the curve
    assert ">0.90<" not in svg                                                 # its meaning goes in the figure's caption
    assert f">{threshold(scores, verdicts, 0.9):.2f}<" in svg


def test_colour_blind_switch(page, one_species):
    page.click("#generate")
    standard = page.locator(".plot-canvas svg g[fill='#0d9488']")
    assert standard.get_attribute("fill-opacity") == "0.5"
    page.check("#colour-blind")
    page.wait_for_selector(".plot-canvas svg g[fill='#0072B2']")
    for colour in ["#0072B2", "#D55E00"]:
        g = page.locator(f".plot-canvas svg g[fill='{colour}']")
        assert g.get_attribute("fill-opacity") is None and g.get_attribute("stroke") == "#ffffff"      # solid, with a thin ring
        assert g.locator("circle").count() > 0
    svg = svg_of(page)
    assert not any(old in svg for old in ["#0d9488", "#c2410c", "#115e59"])    # nothing keeps the standard colours
    page.reload()
    page.wait_for_selector(".plot-canvas svg g[fill='#0072B2']")               # the choice is remembered
    assert page.is_checked("#colour-blind")
    page.uncheck("#colour-blind")
    page.wait_for_selector(".plot-canvas svg g[fill='#0d9488']")


@pytest.mark.parametrize("align, anchor, at", [("left", None, 0.0), ("center", "middle", 0.5), ("right", "end", 1.0)])
def test_title_position(page, one_species, align, anchor, at):
    page.click("#generate")
    page.locator("#title-list input").nth(1).fill("A subtitle")
    page.check(f'input[name="title-align"][value="{align}"]')
    page.wait_for_function(f"document.querySelector('.plot-canvas svg text[font-style]').getAttribute('text-anchor') === {repr(anchor) if anchor else 'null'}")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('A subtitle')")
    x = 64 + at * 660
    for selector in [".plot-canvas svg text[font-style]", '.plot-canvas svg text:text("A subtitle")', '.plot-canvas svg text:text-matches("\\\\d+ positive")']:
        text = page.locator(selector).first
        assert text.get_attribute("text-anchor") == anchor and float(text.get_attribute("x")) == pytest.approx(x, abs=0.1)
    first_key = float(page.locator(".plot-canvas svg circle[fill-opacity='0.6']").first.get_attribute("cx"))
    assert (first_key < 80) == (align == "left")                               # the legend follows


def test_title_position_in_a_grid(page, tmp_path):
    for i, name in enumerate(["Turdus merula", "Parus major"]):
        make_clips(tmp_path / str(i), n=60, seed=i + 1)
        add_species(page, i, name, tmp_path / str(i))
    page.check('input[name="layout"][value="grid"]')
    page.fill("#figure-title", "Spring")
    page.check('input[name="title-align"][value="right"]')
    page.click("#generate")
    figure_title = page.locator('.plot-canvas svg text:text("Spring")').first
    assert figure_title.get_attribute("text-anchor") == "end" and float(figure_title.get_attribute("x")) == pytest.approx(54 + 2 * 330 + 30)
    second = page.locator('.plot-canvas svg text:text("Parus major")').first
    assert second.get_attribute("text-anchor") == "end" and float(second.get_attribute("x")) == pytest.approx(54 + 330 + 30 + 330)


def test_what_is_at_the_top_of_the_plot_is_optional(page, one_species):
    height = lambda: page.locator(".plot-canvas svg").evaluate("(s) => s.viewBox.baseVal.height")
    page.click("#generate")
    full = height()
    svg = svg_of(page)
    assert all(t in svg for t in [">Positive<", ">Negative<", ">Logistic regression<", ">95% confidence band<", ">Threshold<"])
    assert re.search(r">\d+ positive · \d+ negative<", svg)
    assert page.is_hidden("#legend-observed-field")

    page.uncheck("#show-counts")
    page.wait_for_function(f"document.querySelector('.plot-canvas svg').viewBox.baseVal.height < {full}")
    assert not re.search(r">\d+ positive", svg_of(page))
    assert re.search(r"\d+ positive · \d+ negative", svg_of(page))             # still in the description for screen readers
    without_counts = height()

    for switch, gone in [("#legend-points", ">Positive<"), ("#legend-curve", ">Logistic regression<"),
                         ("#legend-band", ">95% confidence band<"), ("#legend-threshold", ">Threshold<")]:
        page.uncheck(switch)
        page.wait_for_function(f"!document.querySelector('.plot-canvas svg').outerHTML.includes('{gone}')")
    assert height() < without_counts                                           # an empty legend takes no room

    # Only the entries go: the marks themselves stay
    assert page.locator('.plot-canvas svg path[fill-opacity="0.14"]').count() == 1
    assert page.locator(".plot-canvas svg path[stroke-dasharray]").count() == 1
    assert page.locator(".plot-canvas svg g[fill='#0d9488'] circle").count() > 0

    page.reload()
    page.wait_for_selector(".plot-canvas svg")
    assert not page.is_checked("#show-counts") and ">Threshold<" not in svg_of(page)


def test_legend_switches_follow_what_is_drawn(page, one_species):
    page.click("#generate")
    page.click(".panel details.advanced summary")
    page.check("#observed")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('Observed share')")
    page.uncheck("#legend-observed")
    page.wait_for_function("!document.querySelector('.plot-canvas svg').outerHTML.includes('Observed share')")
    page.uncheck("#band")
    assert page.is_hidden("#legend-band-field")
    page.wait_for_function("!document.querySelector('.plot-canvas svg').outerHTML.includes('95% confidence band')")


def two_decimals(x):
    """Rounded as people do and as browsers show it: Python's format would turn 0.125 into 0.12."""
    return f"{math.floor(x * 100 + 0.5) / 100:.2f}"


def precision_rows(page, i=0):
    """A precision table on the page, as rows of texts, with its header first."""
    table = page.locator(".precision-table").nth(i)
    return [table.locator("th").all_inner_texts(), *table.locator("tbody tr").evaluate_all("(rows) => rows.map((r) => [...r.cells].map((c) => c.textContent))")]


def wait_for_rows(page, count):
    page.wait_for_function("(n) => document.querySelectorAll('.precision-table tbody tr').length === n", arg=count)


def test_precision_table_is_optional_and_right_by_default(page, one_species):
    scores, verdicts = one_species
    page.click("#generate")
    image = svg_of(page)
    assert not page.is_checked("#precision-table") and page.is_hidden("#precision-options") and page.is_hidden("#download-precision")
    assert page.locator(".precision-table").count() == 0
    page.check("#precision-table")                            # at once, with no click on Generate
    page.wait_for_selector(".precision-table")
    assert page.is_visible("#precision-options") and page.is_visible("#download-precision")
    low, high = verdicts[scores < 0.5], verdicts[scores >= 0.5]
    expected = [["Confidence score", "True positives", "False positives", "Precision"],
                ["< 0.50", str(low.sum()), str(len(low) - low.sum()), two_decimals(low.mean())],
                ["≥ 0.50", str(high.sum()), str(len(high) - high.sum()), two_decimals(high.mean())],
                ["Total", str(verdicts.sum()), str(len(verdicts) - verdicts.sum()), two_decimals(verdicts.mean())]]
    assert precision_rows(page) == expected
    # On the page, right under the image, and not in it
    assert svg_of(page) == image and "True positives" not in image
    assert page.evaluate("document.querySelector('.plot').nextElementSibling.querySelector('.precision-table') !== null")
    page.reload()
    page.wait_for_selector(".precision-table")
    assert page.is_checked("#precision-table") and precision_rows(page) == expected
    page.uncheck("#precision-table")
    page.wait_for_selector(".precision-table", state="detached")
    assert svg_of(page) == image


def test_precision_table_with_scores_chosen_by_the_user(page, one_species):
    scores, verdicts = one_species
    page.check("#precision-table")
    page.click("#generate")
    page.fill("#cut-new", "0,3")
    page.press("#cut-new", "Enter")
    page.fill("#cut-new", "0.8")
    page.click("#cut-add")
    wait_for_rows(page, 5)
    body = precision_rows(page)[1:]
    assert [row[0] for row in body] == ["< 0.30", "≥ 0.30 and < 0.50", "≥ 0.50 and < 0.80", "≥ 0.80", "Total"]
    middle = verdicts[(scores >= 0.5) & (scores < 0.8)]
    assert body[2] == ["≥ 0.50 and < 0.80", str(middle.sum()), str(len(middle) - middle.sum()), two_decimals(middle.mean())]
    raw, name = download(page, "#download-precision")
    rows = read_csv(raw)
    assert name == "birdscore_precision.csv" and len(rows) == 5
    assert rows[2]["range"] == "0.5 to below 0.8" and (rows[2]["score_from"], rows[2]["score_to"]) == ("0.5", "0.8")
    assert int(rows[2]["true_positives"]) == middle.sum() and float(rows[2]["precision"]) == pytest.approx(middle.mean(), abs=1e-4)
    assert rows[4]["range"] == "total" and int(rows[4]["false_positives"]) == len(verdicts) - verdicts.sum()
    assert rows[0]["species"] == "Turdus_merula"

    for value in ["0.3", "0.5", "0.8"]:                       # only the total is left, and a chosen score goes when turned off
        page.click(f'#cut-chips [data-value="{value}"]')
    wait_for_rows(page, 1)
    assert precision_rows(page)[1][0] == "Total"
    assert page.locator('#cut-chips [data-value="0.3"]').count() == 0 and page.locator('#cut-chips [data-value="0.5"]').count() == 1


def test_a_bad_precision_score_says_how_to_fix_it(page):
    page.check("#precision-table")
    page.fill("#cut-new", "1")
    page.click("#cut-add")
    assert page.is_visible("#cut-error") and "above 0 and below 1" in page.inner_text("#cut-error")
    page.fill("#cut-new", "0.7")
    page.click("#cut-add")
    assert page.is_hidden("#cut-error") and page.locator('#cut-chips [data-value="0.7"]').get_attribute("aria-pressed") == "true"


@pytest.mark.parametrize("lang, head, between", [
    ("es", ["Puntuación de confianza", "Verdaderos positivos", "Falsos positivos", "Precisión"], "≥ 0,30 y < 0,50"),
    ("ca", ["Puntuació de confiança", "Veritables positius", "Falsos positius", "Precisió"], "≥ 0,30 i < 0,50"),
])
def test_precision_table_follows_the_language_of_the_site(page, one_species, lang, head, between):
    page.check("#precision-table")
    page.fill("#cut-new", "0.3")
    page.click("#cut-add")
    page.click("#generate")
    page.select_option("#lang", lang)
    page.wait_for_function("(h) => document.querySelector('.precision-table th').textContent === h", arg=head[0])
    rows = precision_rows(page)
    assert rows[0] == head and rows[2][0] == between
    page.select_option("#plot-lang", "en")                    # the plot's own language is for the image
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('Logistic regression')")
    assert precision_rows(page)[0] == head


def test_precision_tables_with_one_image_per_species(page, tmp_path):
    names = ["Turdus merula", "Parus major"]
    for i, name in enumerate(names):
        make_clips(tmp_path / str(i), n=60, seed=i + 1)
        add_species(page, i, name, tmp_path / str(i))
    page.check("#precision-table")
    page.click("#generate")
    assert page.locator(".plot").count() == 2 and page.locator(".precision-table").count() == 2
    for i, name in enumerate(names):
        assert precision_rows(page, i)[0][0] == "Confidence score"        # one species, so no column of species
        assert page.locator(".precision h3").nth(i).inner_text().endswith(name)
        assert page.evaluate("(i) => document.querySelectorAll('.plot')[i].nextElementSibling.querySelector('.precision-table') === document.querySelectorAll('.precision-table')[i]", i)


def test_precision_table_under_a_grid_leaves_the_files_alone(page, tmp_path):
    names = ["Turdus merula", "Parus major", "Erithacus rubecula"]
    for i, name in enumerate(names):
        make_clips(tmp_path / str(i), n=60, seed=i + 1)
        add_species(page, i, name, tmp_path / str(i))
    page.check('input[name="layout"][value="grid"]')
    page.fill("#rows", "2")
    page.fill("#cols", "2")
    page.click("#generate")
    image = svg_of(page)
    png = png_info(download(page, "#download-plots")[0])[:2]
    page.check("#precision-table")
    page.wait_for_selector(".precision-table")
    assert svg_of(page) == image and page.locator(".precision-table").count() == 1
    rows = precision_rows(page)
    assert rows[0] == ["Species", "Confidence score", "True positives", "False positives", "Precision"] and len(rows) == 10
    assert [rows[i][0] for i in (1, 4, 7)] == names and [rows[i][1] for i in (3, 6, 9)] == ["Total"] * 3
    assert png_info(download(page, "#download-plots")[0])[:2] == png            # the table does not go in the image
    page.select_option("#format", "svg")
    page.wait_for_selector("#dpi-field", state="hidden")
    assert "True positives" not in download(page, "#download-plots")[0].decode()
    csv_rows = read_csv(download(page, "#download-precision")[0])
    assert sorted({r["species"] for r in csv_rows}) == sorted(names) and len(csv_rows) == 9


def test_precision_table_when_no_curve_can_be_fitted(page, tmp_path):
    make_clips(tmp_path / "a", n=30)
    page.locator(".species-name").fill("Turdus merula")
    browse(page, 0, "pos", tmp_path / "a" / "Positive")
    page.check("#precision-table")
    page.click("#generate")
    assert "No curve can be fitted" in svg_of(page)
    assert precision_rows(page)[-1][-1] == "1.00"             # with no negatives, the total's precision is 1


def test_a_precision_table_with_many_digits_does_not_widen_the_page(browser, site, tmp_path):
    context = browser.new_context(viewport={"width": 320, "height": 800}, locale="ca-ES")
    page = context.new_page()
    page.goto(site)
    make_clips(tmp_path / "a", n=60)
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.check("#precision-table")
    page.click("#generate")
    page.wait_for_selector(".precision-table")
    assert page.evaluate("document.documentElement.scrollWidth") <= 320
    context.close()


def test_scale_switch_matches_a_raw_fit(page, one_species):
    scores, verdicts = one_species
    page.click("#generate")
    page.click(".panel details.advanced summary")
    page.select_option("#scale", "raw")
    page.wait_for_function("document.querySelector('#methods').textContent.includes('confidence score itself')")
    assert page.locator("#results-table tbody td").nth(4).inner_text().startswith(f"{threshold(scores, verdicts, 0.9, logit=False):.2f}")
    results = read_csv(download(page, "#download-results")[0])
    assert results[0]["score_scale"] == "raw" and float(results[0]["aic_raw"]) > 0 and float(results[0]["aic_logit"]) > 0


# ── Images ───────────────────────────────────────────────────────────────────────────────────────────────
def png_info(raw):
    assert raw[:8] == b"\x89PNG\r\n\x1a\n"
    width, height = struct.unpack(">II", raw[16:24])
    at, chunks = 8, []
    while at < len(raw):
        size = struct.unpack(">I", raw[at:at + 4])[0]
        kind = raw[at + 4:at + 8]
        chunks.append((kind, raw[at + 8:at + 8 + size]))
        at += size + 12
    return width, height, chunks


@pytest.mark.parametrize("dpi", [150, 300, 600])
def test_png_says_its_resolution(page, one_species, dpi):
    page.click("#generate")
    page.select_option("#dpi", str(dpi))
    raw, name = download(page, "#download-plots")
    assert name == "birdscore_Turdus_merula.png"
    width, height, chunks = png_info(raw)
    svg_w, svg_h = page.locator(".plot-canvas svg").evaluate("(s) => [s.viewBox.baseVal.width, s.viewBox.baseVal.height]")
    assert width == round(svg_w * dpi / 96) and height == round(svg_h * dpi / 96)
    phys = [c for c in chunks if c[0] == b"pHYs"]
    assert len(phys) == 1 and chunks[1][0] == b"pHYs"
    ppm, _, unit = struct.unpack(">IIB", phys[0][1])
    assert unit == 1 and ppm == round(dpi / 0.0254)


def test_svg_file_is_standalone(page, one_species):
    page.click("#generate")
    page.select_option("#format", "svg")
    page.wait_for_selector("#dpi-field", state="hidden")
    raw, name = download(page, "#download-plots")
    text = raw.decode()
    assert name.endswith(".svg") and text.startswith("<svg xmlns=") and "http://www.w3.org/2000/svg" in text
    import xml.etree.ElementTree as ET
    assert ET.fromstring(text).tag.endswith("svg")


def test_one_image_per_species_gives_a_zip(page, tmp_path):
    make_clips(tmp_path / "a", seed=1)
    make_clips(tmp_path / "b", seed=2)
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    add_species(page, 1, "Parus major", tmp_path / "b")
    page.click("#generate")
    assert page.locator(".plot-canvas svg").count() == 2
    assert page.inner_text("#download-plots-text") == "Download all (.zip)"
    raw, name = download(page, "#download-plots")
    z = zipfile.ZipFile(io.BytesIO(raw))
    assert name == "birdscore_plots.zip" and z.testzip() is None
    assert sorted(z.namelist()) == ["birdscore_Parus_major.png", "birdscore_Turdus_merula.png"]
    assert png_info(z.read("birdscore_Parus_major.png"))[0] > 2000


def test_one_image_with_a_grid(page, tmp_path):
    for i, name in enumerate(["Turdus merula", "Parus major", "Erithacus rubecula", "Sylvia atricapilla"]):
        make_clips(tmp_path / str(i), n=60, seed=i + 1)
        add_species(page, i, name, tmp_path / str(i))
    page.click("#generate")
    single = page.locator(".plot-canvas svg").nth(0).evaluate("(s) => s.viewBox.baseVal.width")
    page.check('input[name="layout"][value="grid"]')
    page.wait_for_function("document.querySelectorAll('.plot-canvas svg').length === 1")
    assert page.is_visible("#grid-size")
    assert page.get_attribute("#rows", "placeholder") == "2" and page.get_attribute("#cols", "placeholder") == "2"
    size = lambda: page.locator(".plot-canvas svg").evaluate("(s) => [s.viewBox.baseVal.width, s.viewBox.baseVal.height]")
    w2, h2 = size()
    assert w2 == pytest.approx(54 + 2 * 330 + 30 + 20)
    page.fill("#rows", "2")
    page.fill("#cols", "3")                                       # the user's example: a 2 x 3 grid
    page.wait_for_function(f"document.querySelector('.plot-canvas svg').viewBox.baseVal.width > {w2}")
    assert size()[0] == pytest.approx(54 + 3 * 330 + 2 * 30 + 20)
    page.fill("#rows", "1")
    page.fill("#cols", "2")                                       # too small: it grows, and says so
    page.wait_for_selector("#grid-note:not([hidden])")
    assert "2 rows" in page.inner_text("#grid-note")
    page.fill("#figure-title", "Spring validation")
    page.fill("#figure-subtitle", "All sites")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('All sites')")       # the last change
    svg = svg_of(page)
    assert "Spring validation" in svg and all(n in svg for n in ["Turdus merula", "Parus major", "Erithacus rubecula", "Sylvia atricapilla"])
    assert svg.count(">True positive rate<") == 2 and svg.count(">Confidence score<") == 2  # shared along each row and column
    raw, name = download(page, "#download-plots")
    assert name == "birdscore_all_species.png" and png_info(raw)[0] > single


# ── Language ─────────────────────────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("lang, title, curve, counts, positive, axis", [
    ("es", "Tus detecciones", "Regresión logística", r"\d+ positivos · \d+ negativos", ">Positivo<", "Puntuación de confianza"),
    ("ca", "Les teves deteccions", "Regressió logística", r"\d+ positius · \d+ negatius", ">Positiu<", "Puntuació de confiança"),
])
def test_other_languages_keep_the_axis_names_and_translate_the_rest(page, one_species, lang, title, curve, counts, positive, axis):
    page.click("#generate")
    page.select_option("#lang", lang)
    assert page.inner_text("#data-title") == title and page.get_attribute("html", "lang") == lang
    page.wait_for_function(f"document.querySelector('.plot-canvas svg').outerHTML.includes('{curve}')")
    svg = svg_of(page)
    assert ">Confidence score<" in svg and ">True positive rate<" in svg       # as the owner asked
    assert page.input_value("#x-label") == "Confidence score" and page.input_value("#y-label") == "True positive rate"
    assert positive in svg and re.search(counts, svg)
    assert ">0,2<" in svg and re.search(r">0,\d\d<", svg)                      # decimal commas on the axis and the threshold
    page.fill("#x-label", axis)
    page.wait_for_function(f"document.querySelector('.plot-canvas svg').outerHTML.includes('{axis}')")


@pytest.mark.parametrize("browser_locale, expected", [("ca-ES", "ca"), ("es-ES", "es"), ("en-GB", "en"), ("fr-FR", "en")])
def test_the_first_visit_follows_the_browser_language(browser, site, browser_locale, expected):
    context = browser.new_context(locale=browser_locale)
    page = context.new_page()
    page.goto(site)
    assert page.get_attribute("html", "lang") == expected and page.input_value("#lang") == expected
    assert page.input_value("#plot-lang") == expected
    context.close()


def test_the_language_is_remembered(page):
    page.select_option("#lang", "ca")
    page.reload()
    assert page.input_value("#lang") == "ca" and page.inner_text("#data-title") == "Les teves deteccions"


def test_the_plot_language_can_differ_from_the_site(page, one_species):
    page.select_option("#lang", "es")
    assert page.input_value("#plot-lang") == "es"
    page.click("#generate")
    page.select_option("#plot-lang", "en")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('Logistic regression')")
    assert page.inner_text("#data-title") == "Tus detecciones" and re.search(r">0\.\d\d<", svg_of(page))


def test_singular_in_both_languages(page, tmp_path):
    root = tmp_path / "a"
    (root / "Positive").mkdir(parents=True)
    (root / "Negative").mkdir(parents=True)
    (root / "Positive" / "0.800_1_r_0.0s_3.0s.wav").touch()
    (root / "Positive" / "0.300_1_r_3.0s_6.0s.wav").touch()
    (root / "Negative" / "0.500_1_r_6.0s_9.0s.wav").touch()
    (root / "Negative" / "bad.wav").touch()
    add_species(page, 0, "Parus major", root)
    assert page.locator('.slot[data-part="neg"] .slot-count').inner_text() == "1 clip, score 0.500"
    assert page.locator('.slot[data-part="neg"] .issue summary').inner_text() == "1 file not used"
    page.click("#generate")
    assert page.inner_text("#results-summary") == "One species, 3 clips"
    assert "Only 1 negative clip:" in page.inner_text("#notes") and ">2 positive · 1 negative<" in svg_of(page)
    page.select_option("#lang", "es")
    page.wait_for_function("document.querySelector('.plot-canvas svg').outerHTML.includes('1 negativo<')")
    assert page.inner_text("#results-summary") == "Una especie, 3 clips"
    assert "Solo 1 clip negativo:" in page.inner_text("#notes") and ">2 positivos · 1 negativo<" in svg_of(page)
    assert page.locator('.slot[data-part="neg"] .slot-count').inner_text() == "1 clip, puntuación 0,500"


def test_every_text_has_its_translations(page):
    keys = set(page.evaluate("""() => {
      const doc = [...document.querySelectorAll('[data-i18n]')].map((e) => e.dataset.en);
      const tpl = [...document.querySelector('#species-template').content.querySelectorAll('[data-i18n]')].map((e) => e.dataset.en ?? e.textContent.replace(/\\s+/g, ' ').trim());
      const attrs = [...document.querySelectorAll('[data-i18n-attr]'), ...document.querySelector('#species-template').content.querySelectorAll('[data-i18n-attr]')]
        .flatMap((e) => e.dataset.i18nAttr.split(' ').map((name) => e.dataset[`en_${name.replace(/-/g, '_')}`] ?? e.getAttribute(name)));
      return [...doc, ...tpl, ...attrs];
    }"""))
    source = (APP / "app.js").read_text(encoding="utf-8")
    keys |= set(re.findall(r"""(?<![\w.])(?:t|T)\('((?:[^'\\]|\\.)*)'""", source))
    keys |= set(re.findall(r"""tIn\(lang, '((?:[^'\\]|\\.)*)'""", source))
    for match in re.findall(r"""t\(\w+(?:\.\w+)? === '\w+' \? '((?:[^'\\]|\\.)*)' : '((?:[^'\\]|\\.)*)'""", source):
        keys |= set(match)
    for match in re.findall(r"""tn\([^,]+, '((?:[^'\\]|\\.)*)', '((?:[^'\\]|\\.)*)'""", source):
        keys |= set(match)
    for match in re.findall(r"""\[a\.n\w+, '([^']+)', '([^']+)'\]""", source):
        keys |= set(match)
    keys |= set(re.findall(r"""(?:no_score|score_range|not_audio): '([^']+)'""", source))
    keys |= {"Positive", "Negative", "Ignore"}
    es, ca = page.evaluate("[Object.keys(I18n.ES), Object.keys(I18n.CA)]")
    assert not sorted(k for k in keys if k not in es), "missing in Spanish"
    assert not sorted(k for k in keys if k not in ca), "missing in Catalan"
    assert sorted(es) == sorted(ca)                       # neither language holds a text the other lacks
    assert len(keys) > 100


# ── Tables ───────────────────────────────────────────────────────────────────────────────────────────────
def open_table(page, path):
    with page.expect_file_chooser() as chooser:
        page.click("#choose-table")
    chooser.value.set_files(str(path))
    page.wait_for_selector("#table-setup:not([hidden])")


def table_rows():
    rng = np.random.default_rng(5)
    rows = []
    for species in ["Turdus merula", "Parus major"]:
        for i in range(80):
            s = round(float(rng.uniform(0.1, 1.0)), 3)
            rows.append((species, s, bool(rng.random() < s), f"{s:.3f}_{i + 1}_SITE02_20240503_0630{i % 60:02d}_{i}.0s_{i + 3}.0s.wav"))
    return rows


def test_csv_with_semicolons_and_decimal_commas(page, tmp_path):
    rows = table_rows()
    path = tmp_path / "detections.csv"
    lines = ["Especie;Confianza;Resultado;Archivo;Observador"] + [f"{a};{str(b).replace('.', ',')};{'sí' if c else 'no'};{d};Obs1" for a, b, c, d in rows]
    path.write_bytes(("\n".join(lines)).encode("cp1252"))
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    assert page.input_value('select[data-role="species"]') == "0" and page.input_value('select[data-role="score"]') == "1"
    assert page.input_value('select[data-role="result"]') == "2" and page.input_value('select[data-role="file"]') == "3"
    pos = sum(r[2] for r in rows)
    assert f"{len(rows)} detections: {pos} positive, {len(rows) - pos} negative." == page.inner_text("#table-status")
    page.click("#generate")
    assert page.locator(".plot-canvas svg").count() == 2
    out = read_csv(download(page, "#download-data")[0])
    assert len(out) == len(rows) and out[0]["Observador"] == "Obs1" and out[0]["date"] == "2024-05-03"
    assert out[0]["site_or_recorder"] == "SITE02" and float(out[0]["confidence_score"]) == rows[0][1]


def test_a_saved_csv_can_be_read_back(page, one_species, tmp_path):
    page.click("#generate")
    for style, sep in [("comma", ","), ("semicolon", ";")]:
        page.select_option("#csv-format", style)
        raw, _ = download(page, "#download-data")
        path = tmp_path / f"{style}.csv"
        path.write_bytes(raw)
        before = read_csv(raw, sep)
        page.check('input[name="source"][value="table"]')
        open_table(page, path)
        assert page.input_value('select[data-role="score"]') == "4" and page.input_value('select[data-role="result"]') == "2"
        assert page.inner_text("#table-status").startswith(f"{len(before)} detections")
        page.check('input[name="source"][value="folders"]')


def test_excel_with_a_title_line_dates_and_several_sheets(page, tmp_path):
    rows = table_rows()
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.active.append(["Nothing here"])
    sheet = book.create_sheet("Checked")
    sheet.append(["Validation of spring"])
    sheet.append(["Scientific name", "Confidence", "Validated", "Date", "Time", "Recorder"])
    for a, b, c, _ in rows:
        sheet.append([a, b, c, datetime.datetime(2024, 5, 3), datetime.time(21, 15, 30), "R7"])
    path = tmp_path / "checked.xlsx"
    book.save(path)
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    assert page.is_visible("#sheet-field") and page.input_value("#sheet") == "1"
    assert page.locator("#table-preview th").all_inner_texts()[:2] == ["Scientific name", "Confidence"]
    assert [page.input_value(f'select[data-role="{r}"]') for r in ["species", "score", "result", "date", "time", "site"]] == ["0", "1", "2", "3", "4", "5"]
    page.click("#generate")
    out = read_csv(download(page, "#download-data")[0])
    assert len(out) == len(rows)
    assert (out[0]["date"], out[0]["time"], out[0]["site_or_recorder"]) == ("2024-05-03", "21:15:30", "R7")
    assert float(out[0]["confidence_score"]) == rows[0][1]                     # numbers keep every decimal
    assert sum(r["result"] == "positive" for r in out) == sum(r[2] for r in rows)


def test_old_excel_format(page, tmp_path):
    book = xlwt.Workbook()
    sheet = book.add_sheet("Hoja1")
    for j, head in enumerate(["species", "score", "tp"]):
        sheet.write(0, j, head)
    for i, (a, b, c, _) in enumerate(table_rows()[:60], start=1):
        sheet.write(i, 0, a)
        sheet.write(i, 1, b)
        sheet.write(i, 2, "TP" if c else "FP")
    path = tmp_path / "old.xls"
    book.save(str(path))
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    assert page.inner_text("#table-status").startswith("60 detections")


def test_one_species_without_a_species_column_and_unknown_values(page, tmp_path):
    path = tmp_path / "plain.csv"
    path.write_text("score,check\n0.2,bad\n0.3,good\n0.5,bad\n0.7,good\n0.9,good\n0.4,unsure\nx,good\n", encoding="utf-8")
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    page.select_option('select[data-role="result"]', "1")
    assert page.input_value('select[data-role="species"]') == "-1" and page.is_visible("#fixed-species-field")
    assert page.is_disabled("#generate")
    assert page.locator("#value-list select").count() == 3
    page.locator("#value-list select").nth(0).select_option("negative")       # bad
    page.locator("#value-list select").nth(1).select_option("positive")       # good
    assert page.is_disabled("#generate")                                      # still no species name
    page.fill("#fixed-species", "Parus major")
    assert page.is_enabled("#generate")
    status = page.inner_text("#table-status")
    assert status.startswith("5 detections: 3 positive, 2 negative.") and "2 rows not used" in status
    page.click("#generate")
    assert "Parus major" in svg_of(page)
    assert "very uncertain" in page.inner_text("#notes")


def test_a_file_that_is_not_a_table_says_so(page, tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("just one line of text", encoding="utf-8")
    page.check('input[name="source"][value="table"]')
    with page.expect_file_chooser() as chooser:
        page.click("#choose-table")
    chooser.value.set_files(str(path))
    page.wait_for_selector("#table-error:not([hidden])")
    assert "could not be read as a table" in page.inner_text("#table-error")


# ── What must not pass silently ──────────────────────────────────────────────────────────────────────────
def test_files_that_cannot_be_used_are_listed(page, tmp_path):
    make_clips(tmp_path / "a", n=40)
    (tmp_path / "a" / "Positive" / "SITE01_20240501_203000_0.0s_3.0s.wav").touch()
    (tmp_path / "a" / "Positive" / "1.7_3_rec.wav").touch()
    (tmp_path / "a" / "Positive" / "notes.txt").touch()
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    details = page.locator('.slot[data-part="pos"] .issue')
    assert details.locator("summary").inner_text() == "3 files not used"
    details.locator("summary").click()
    text = details.inner_text()
    assert "no score at the start" in text and "above 1" in text and "not an audio file" in text


def test_a_swapped_folder_is_flagged(page, tmp_path):
    make_clips(tmp_path / "a", n=60)
    page.locator(".species-name").fill("Turdus merula")
    browse(page, 0, "pos", tmp_path / "a" / "Negative")
    browse(page, 0, "neg", tmp_path / "a" / "Positive")
    assert "Negative" in page.locator('.slot[data-part="pos"] .warn').inner_text()
    page.click("#generate")
    text = page.inner_text("#notes")
    assert "swapped" in text
    assert page.locator(".plot-canvas svg path[stroke-dasharray]").count() == 0    # no threshold from a falling curve


def test_only_positive_clips_gives_points_and_a_note(page, tmp_path):
    make_clips(tmp_path / "a", n=30)
    page.locator(".species-name").fill("Turdus merula")
    browse(page, 0, "pos", tmp_path / "a" / "Positive")
    page.click("#generate")
    assert "no negative clips" in page.inner_text("#notes")
    assert "No curve can be fitted" in svg_of(page)
    assert page.locator(".plot-canvas svg circle").count() > 5


def test_years_that_split_the_verdicts_are_flagged(page, tmp_path):
    root = tmp_path / "a"
    (root / "Positive").mkdir(parents=True)
    (root / "Negative").mkdir(parents=True)
    for i in range(30):
        (root / "Positive" / f"{0.2 + i / 50:.3f}_{i}_R1_20220501_203000_{i}.0s_{i + 3}.0s.wav").touch()
        (root / "Negative" / f"{0.1 + i / 50:.3f}_{i}_R2_20250501_203000_{i}.0s_{i + 3}.0s.wav").touch()
    add_species(page, 0, "Turdus merula", root)
    page.click("#generate")
    text = page.inner_text("#notes")
    assert "from 2022" in text and "from 2025" in text and "the year" in text and "site or recorder" in text


def test_a_name_is_required_and_must_be_unique(page, tmp_path):
    make_clips(tmp_path / "a", n=30)
    browse(page, 0, "pos", tmp_path / "a" / "Positive")
    page.click("#generate")
    assert "Write the species name" in page.inner_text(".name-error")
    assert page.is_hidden("#results")
    assert page.evaluate("document.activeElement.classList.contains('species-name')")
    page.locator(".species-name").fill("Turdus merula")
    assert page.is_hidden(".name-error")
    page.click("#add-species")
    page.locator(".species-name").nth(1).fill("Turdus merula")
    browse(page, 1, "pos", tmp_path / "a" / "Positive")
    page.click("#generate")
    assert "already has this name" in page.locator(".name-error").nth(1).inner_text()


def test_names_from_files_are_never_run_as_code(page, tmp_path):
    make_clips(tmp_path / "a", n=60)
    bad = '<img src=x onerror="window.hacked=1">&<b>'
    add_species(page, 0, bad, tmp_path / "a")
    page.fill("#x-label", '<script>window.hacked=1</script>')
    page.click("#generate")
    page.wait_for_function("document.querySelector('.plot-canvas svg text')")
    assert page.evaluate("window.hacked") is None
    assert page.locator(".plot-canvas svg img, .plot-canvas svg script, #results-table img, #title-list img").count() == 0
    assert bad in page.locator("#results-table tbody td").first.inner_text()
    import xml.etree.ElementTree as ET
    ET.fromstring(svg_of(page))                               # still a valid file


def test_a_large_table_stays_usable(page, tmp_path):
    rng = np.random.default_rng(9)
    path = tmp_path / "large.csv"
    with path.open("w", encoding="utf-8") as f:
        f.write("species,score,result\n")
        for i in range(120000):
            s = rng.uniform(0.1, 1)
            f.write(f"Parus major,{s:.3f},{int(rng.random() < s)}\n")
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    assert page.inner_text("#table-status").startswith("120,000 detections")
    page.click("#generate")
    page.wait_for_selector(".plot-canvas svg")
    assert page.locator(".plot-canvas svg g circle").count() < 10010
    assert "sample of 5,000 points" in page.inner_text("#notes")
    assert len(read_csv(download(page, "#download-data")[0])) == 120000


def test_changing_the_data_marks_the_plot_as_old(page, one_species):
    page.click("#generate")
    assert page.is_hidden("#stale")
    page.locator(".species-name").fill("Another name")
    assert page.is_visible("#stale")
    page.click("#generate")
    assert page.is_hidden("#stale") and "Another name" in svg_of(page)


# ── The last project ─────────────────────────────────────────────────────────────────────────────────────
def test_the_last_plots_come_back_after_a_reload(page, one_species):
    page.fill("#x-label", "My score")
    page.click('#target-chips [data-value="0.95"]')
    page.click("#generate")
    before = svg_of(page)
    page.reload()
    page.wait_for_selector(".plot-canvas svg")
    assert svg_of(page) == before
    assert page.input_value("#x-label") == "My score"
    assert page.locator(".species-name").input_value() == "Turdus_merula"
    assert page.locator('#target-chips [data-value="0.95"]').get_attribute("aria-pressed") == "true"
    page.click("#start-over")
    page.click("#ask-keep")
    assert page.is_visible("#results")
    page.click("#start-over")
    page.click("#ask-remove")
    assert page.is_hidden("#results") and page.is_visible("#placeholder")
    page.reload()
    assert page.is_hidden("#results") and page.locator(".species-name").input_value() == ""


def test_a_table_comes_back_too(page, tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("species,score,result\n" + "".join(f"Parus major,{0.1 + i / 50:.3f},{'yes' if i % 3 else 'no'}\n" for i in range(45)), encoding="utf-8")
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    page.click("#generate")
    page.reload()
    page.wait_for_selector(".plot-canvas svg")
    assert "t.csv" in page.inner_text("#table-name") and page.is_enabled("#generate")


# ── Works from disk, on a phone, by keyboard ─────────────────────────────────────────────────────────────
def test_works_opened_straight_from_disk(browser, tmp_path):
    context = browser.new_context(accept_downloads=True, locale="en-GB")
    page = context.new_page()
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.goto((APP / "index.html").as_uri())
    make_clips(tmp_path / "a", n=60)
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.click("#generate")
    assert png_info(download(page, "#download-plots")[0])[0] > 2000
    book = openpyxl.Workbook()
    book.active.append(["species", "score", "result"])
    for i in range(20):
        book.active.append(["Parus major", 0.1 + i / 25, i % 2])
    book.save(tmp_path / "t.xlsx")
    page.check('input[name="source"][value="table"]')
    open_table(page, tmp_path / "t.xlsx")
    assert page.inner_text("#table-status").startswith("20 detections")
    assert not errors
    context.close()


@pytest.mark.parametrize("width", [320, 390, 768])
def test_no_sideways_scrolling(browser, site, tmp_path, width):
    context = browser.new_context(viewport={"width": width, "height": 800}, locale="en-GB")
    page = context.new_page()
    page.goto(site)
    make_clips(tmp_path / "a", n=60)
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.click("#generate")
    for lang in ["en", "es", "ca"]:                       # the longer texts are what overflows
        page.select_option("#lang", lang)
        page.check('input[name="source"][value="folders"]')
        assert page.evaluate("document.documentElement.scrollWidth") <= width, lang
        page.check('input[name="source"][value="table"]')
        assert page.evaluate("document.documentElement.scrollWidth") <= width, lang
    context.close()


def test_every_control_has_a_name(page, one_species):
    page.click("#generate")
    nameless = page.evaluate("""() => [...document.querySelectorAll('button, input:not([type=file]), select, summary')]
      .filter((e) => e.offsetParent !== null)
      .filter((e) => !(e.getAttribute('aria-label') || e.textContent.trim() || e.labels?.length || e.getAttribute('aria-labelledby')))
      .map((e) => e.outerHTML.slice(0, 80))""")
    assert not nameless, nameless


def test_every_plot_setting_redraws_at_once(page, tmp_path):
    """After the first Generate, every setting that shows on the plot changes it with no further click."""
    for i, name in enumerate(["Turdus merula", "Parus major"]):
        make_clips(tmp_path / str(i), n=80, seed=i + 1)
        add_species(page, i, name, tmp_path / str(i))
    page.click("#generate")
    page.wait_for_selector(".plot-canvas svg")
    page.click(".panel details.advanced summary")

    def add(field, button, value):
        page.fill(field, value)
        page.click(button)

    def amount():
        page.locator("#jitter-amount").evaluate("(e) => { e.value = 0.2; e.dispatchEvent(new Event('input', { bubbles: true })); }")

    settings = {
        "a probability chip": lambda: page.click('#target-chips [data-value="0.95"]'),
        "another probability": lambda: add("#target-new", "#target-add", "0.8"),
        "a score": lambda: add("#score-new", "#score-add", "0.5"),
        "title": lambda: page.locator("#title-list input").nth(0).fill("New title"),
        "subtitle": lambda: page.locator("#title-list input").nth(1).fill("New subtitle"),
        "italics": lambda: page.uncheck("#italic"),
        "title in the centre": lambda: page.check('input[name="title-align"][value="center"]'),
        "title on the right": lambda: page.check('input[name="title-align"][value="right"]'),
        "totals": lambda: page.uncheck("#show-counts"),
        "legend, points": lambda: page.uncheck("#legend-points"),
        "legend, curve": lambda: page.uncheck("#legend-curve"),
        "legend, band": lambda: page.uncheck("#legend-band"),
        "legend, threshold": lambda: page.uncheck("#legend-threshold"),
        "x axis": lambda: page.fill("#x-label", "X name"),
        "y axis": lambda: page.fill("#y-label", "Y name"),
        "jitter amount": amount,
        "jitter": lambda: page.uncheck("#jitter"),
        "colour-blind colours": lambda: page.check("#colour-blind"),
        "plot language": lambda: page.select_option("#plot-lang", "ca"),
        "score scale": lambda: page.select_option("#scale", "raw"),
        "band": lambda: page.uncheck("#band"),
        "observed share": lambda: page.check("#observed"),
        "legend, observed share": lambda: page.uncheck("#legend-observed"),
        "one image with all species": lambda: page.check('input[name="layout"][value="grid"]'),
        "columns": lambda: page.fill("#cols", "1"),
        "rows": lambda: page.fill("#rows", "3"),
        "image title": lambda: page.fill("#figure-title", "Image title"),
        "image subtitle": lambda: page.fill("#figure-subtitle", "Image subtitle"),
    }
    for name, change in settings.items():
        before = svg_of(page)
        change()
        try:
            page.wait_for_function("(b) => document.querySelector('.plot-canvas svg').outerHTML !== b", arg=before, timeout=3000)
        except Exception:
            pytest.fail(f"{name} did not change the plot")
    assert page.locator("#generate").is_enabled()


def test_the_page_links_are_stamped_when_published():
    html = (APP / "index.html").read_text(encoding="utf-8")
    links = re.findall(r'(?:src|href)="([\w./-]+\.(?:js|css)[^"]*)"', html)
    assert len(links) == 5 and all(link.endswith("?v=dev") for link in links), links
    workflow = (APP.parent / ".github" / "workflows" / "pages.yml").read_text(encoding="utf-8")
    assert 's/?v=dev/?v=${GITHUB_SHA::8}/g' in workflow


def test_a_new_version_never_runs_old_scripts(browser, tmp_path):
    """Pages lets browsers keep each file for 10 minutes, so a reload after a push can get the new page and old scripts."""
    class Pages(http.server.SimpleHTTPRequestHandler):
        def end_headers(self):
            self.send_header("Cache-Control", "max-age=600")
            super().end_headers()

        def log_message(self, *args):
            pass

    folder = tmp_path / "site"
    shutil.copytree(APP, folder)

    def publish(build, stamp_time):
        page_file, script = folder / "index.html", folder / "app.js"
        page_file.write_text((APP / "index.html").read_text(encoding="utf-8").replace("?v=dev", f"?v={build}"), encoding="utf-8")
        script.write_text((APP / "app.js").read_text(encoding="utf-8") + f"\nwindow.BUILD = '{build}';\n", encoding="utf-8")
        for path in (page_file, script):
            os.utime(path, (stamp_time, stamp_time))

    now = int(os.path.getmtime(APP / "index.html"))
    publish("one", now - 7200)
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), partial(Pages, directory=str(folder)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    context = browser.new_context()
    page = context.new_page()
    page.goto(f"http://127.0.0.1:{server.server_port}/")
    assert page.evaluate("window.BUILD") == "one"
    publish("two", now - 3600)
    page.reload()
    assert page.evaluate("window.BUILD") == "two"
    context.close()
    server.shutdown()


def test_nothing_is_sent_anywhere(browser, site, tmp_path):
    """Only the site's own files are requested, whatever the user does."""
    context = browser.new_context(accept_downloads=True)
    page = context.new_page()
    seen = []
    page.on("request", lambda r: seen.append((r.method, r.url)))
    page.goto(site)
    make_clips(tmp_path / "a", n=60)
    add_species(page, 0, "Turdus merula", tmp_path / "a")
    page.click("#generate")
    download(page, "#download-plots")
    download(page, "#download-data")
    path = tmp_path / "t.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["species", "score", "result"])
    book.active.append(["Parus major", 0.5, 1])
    book.save(path)
    page.check('input[name="source"][value="table"]')
    open_table(page, path)
    others = [u for m, u in seen if not (u.startswith(site) or u.startswith("blob:") or u.startswith("data:"))]
    assert not others and all(m == "GET" for m, _ in seen)
    context.close()
